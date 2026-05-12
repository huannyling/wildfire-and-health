from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path("/Users/wenhuan/Documents/Project_Claude/wildfire-lit-review")
FINAL_DIR = ROOT / "outputs" / "final"

SOURCE_MANUSCRIPT_DOCX = FINAL_DIR / "main_manuscript wildfire 2026-05.docx"
SOURCE_MANUSCRIPT_MD = FINAL_DIR / "main_manuscript.md"
SOURCE_COVER_MD = FINAL_DIR / "nature_cover_letter.md"

OUTPUT_MANUSCRIPT_DOCX = FINAL_DIR / "main_manuscript wildfire 2026-05 nature_article.docx"
OUTPUT_COVER_DOCX = FINAL_DIR / "nature_cover_letter 2026-05.docx"


OLD_TO_NEW = {
    "Wildfire activity is increasing under climate change, yet it remains unclear whose health risks are being studied—and whose are not. We analysed 1,999 wildfire-related papers (2000–2024) using a large language model to extract health outcomes, populations, geographies, and causal pathways. We identify four structural gaps. First, reproductive and neurological outcomes are largely absent, while respiratory impacts dominate the literature. Second, the most vulnerable populations—including unhoused individuals, Indigenous communities, and pregnant women—are consistently underrepresented. Third, research effort is highly uneven: the United States produces more than ten times its proportional share of wildfire health studies, whereas many high-risk countries have none. Fourth, across the literature, PM2.5 forms the central pathway linking fire to health, with little attention to social determinants. These patterns reflect systemic biases in study design, funding, and data availability. Together, they suggest that current research priorities do not align with global health needs, limiting the evidence base for protecting the populations most at risk.": "Wildfire activity is rising under climate change, but it remains unclear whether the health literature reflects the populations, outcomes and geographies at greatest risk. Here we use a validated large language model pipeline to extract structured outcome, population, geographic and causal-pathway data from 1,999 wildfire-related papers published between 2000 and 2024. We show that the field is organized around a narrow evidentiary core: respiratory outcomes and PM2.5-mediated pathways dominate, whereas reproductive and neurological outcomes remain sparsely studied despite high public-health relevance. The literature is also poorly aligned with vulnerability: unhoused people, Indigenous communities and pregnant women are consistently underrepresented, and research output is heavily concentrated in a small set of high-income countries, particularly the United States, while many fire-prone countries have little or no identified health evidence. Finally, causal knowledge networks are dense for particulate pathways but largely fail to capture social determinants of exposure and harm. Together, these findings indicate that wildfire health research is shaped by structural biases in study design, infrastructure and funding, limiting its capacity to guide protection of the populations most at risk.",
    "A rapidly growing but geographically concentrated corpus": "Growth and concentration",
    "Neurological and reproductive outcomes are structurally absent from the evidence base": "Reproductive and neurological gaps",
    "Four of eight vulnerable population groups are critically underrepresented": "Missing vulnerable populations",
    "The United States captures 16 times its proportional share of wildfire health research": "Global research-risk mismatch",
    "The wildfire health knowledge network is dominated by a single mechanistic hub": "A PM2.5-centred knowledge network",
    "Research trajectories: mental health emerging, cardiovascular declining, neurological persistently neglected": "Uneven research trajectories",
}


def replace_paragraph_text(doc: Document, old_to_new: dict[str, str]) -> int:
    replaced = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if text in old_to_new:
            para.text = old_to_new[text]
            replaced += 1
    return replaced


def build_manuscript_docx() -> None:
    shutil.copy2(SOURCE_MANUSCRIPT_DOCX, OUTPUT_MANUSCRIPT_DOCX)
    doc = Document(str(OUTPUT_MANUSCRIPT_DOCX))
    replaced = replace_paragraph_text(doc, OLD_TO_NEW)
    expected = len(OLD_TO_NEW)
    if replaced != expected:
        raise RuntimeError(f"Expected {expected} replacements, got {replaced}")
    doc.save(str(OUTPUT_MANUSCRIPT_DOCX))


def build_cover_letter_docx() -> None:
    text = SOURCE_COVER_MD.read_text(encoding="utf-8").strip()
    paragraphs = [block.strip() for block in text.split("\n\n") if block.strip()]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for block in paragraphs:
        cleaned = (
            block.replace("*Nature*", "Nature")
            .replace('**"Wildfire health research fails to protect vulnerable populations,"**', '“Wildfire health research fails to protect vulnerable populations,”')
            .replace("**Article**", "Article")
            .replace("**", "")
            .replace("*", "")
            .replace("  \n", "\n")
        )
        para = doc.add_paragraph()
        para.style = doc.styles["Normal"]
        para.add_run(cleaned)

    doc.save(str(OUTPUT_COVER_DOCX))


def main() -> None:
    if not SOURCE_MANUSCRIPT_DOCX.exists():
        raise FileNotFoundError(SOURCE_MANUSCRIPT_DOCX)
    if not SOURCE_COVER_MD.exists():
        raise FileNotFoundError(SOURCE_COVER_MD)

    build_manuscript_docx()
    build_cover_letter_docx()

    print(OUTPUT_MANUSCRIPT_DOCX)
    print(OUTPUT_COVER_DOCX)


if __name__ == "__main__":
    main()
